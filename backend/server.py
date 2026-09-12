from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import logging
from pathlib import Path
from pydantic import BaseModel, Field, BeforeValidator, ConfigDict
from typing import List, Optional, Annotated, Any
from datetime import datetime, timezone, timedelta
from bson import ObjectId

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import bcrypt
import jwt
import qrcode
import openpyxl
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

# ---------------------------------------------------------------- DB
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------- Models
PyObjectId = Annotated[str, BeforeValidator(str)]


class Record(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    id: Optional[PyObjectId] = Field(default=None, alias="_id")
    kebun: str = ""
    afdeling: str = ""
    blok: str = ""
    code_lsu: str = ""
    luas_ha: float = 0
    jumlah_pokok: float = 0
    titik_sample: str = ""
    koord_x: float = 0
    koord_y: float = 0
    kategori: str = ""
    keterangan: str = ""
    sph: float = 0
    jumlah_pelepah: float = 0
    panjang_pelepah: float = 0
    lebar_petiol: float = 0
    tebal_petiol: float = 0
    panjang_helai_1: float = 0
    panjang_helai_2: float = 0
    lebar_helai_1: float = 0
    lebar_helai_2: float = 0
    jumlah_anak_daun: float = 0
    tanggal_lsu: str = ""
    la: float = 0
    lai: float = 0
    id_actual: str = ""
    created_at: Optional[str] = None


class RecordInput(BaseModel):
    kebun: str = ""
    afdeling: str = ""
    blok: str = ""
    code_lsu: str = ""
    luas_ha: float = 0
    jumlah_pokok: float = 0
    titik_sample: str = ""
    koord_x: float = 0
    koord_y: float = 0
    kategori: str = ""
    keterangan: str = ""
    jumlah_pelepah: float = 0
    panjang_pelepah: float = 0
    lebar_petiol: float = 0
    tebal_petiol: float = 0
    panjang_helai_1: float = 0
    panjang_helai_2: float = 0
    lebar_helai_1: float = 0
    lebar_helai_2: float = 0
    jumlah_anak_daun: float = 0
    tanggal_lsu: str = ""
    la: float = 0
    lai: float = 0


class ImportRow(BaseModel):
    kebun: str = ""
    afdeling: str = ""
    blok: str = ""
    code_lsu: str = ""
    luas_ha: float = 0
    jumlah_pokok: float = 0
    titik_sample: str = ""
    koord_x: float = 0
    koord_y: float = 0
    kategori: str = ""
    keterangan: str = ""
    jumlah_pelepah: float = 0
    panjang_pelepah: float = 0
    lebar_petiol: float = 0
    tebal_petiol: float = 0
    panjang_helai_1: float = 0
    panjang_helai_2: float = 0
    lebar_helai_1: float = 0
    lebar_helai_2: float = 0
    jumlah_anak_daun: float = 0
    tanggal_lsu: str = ""
    la: float = 0
    lai: float = 0


class ImportConfirm(BaseModel):
    rows: List[ImportRow]


class IdList(BaseModel):
    ids: List[str] = []
    size: str = "medium"


class UserRegister(BaseModel):
    email: str
    password: str
    name: str = "User"


class UserLogin(BaseModel):
    email: str
    password: str


# ---------------------------------------------------------------- Auth utils
JWT_ALGORITHM = "HS256"


def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))


def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email,
               "exp": datetime.now(timezone.utc) + timedelta(hours=12), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)


def set_auth_cookie(response: Response, token: str):
    response.set_cookie(key="access_token", value=token, httponly=True,
                        secure=True, samesite="none", max_age=43200, path="/")


async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")


# ---------------------------------------------------------------- QR helpers
def fmt_num(v) -> str:
    """Format koordinat memakai koma sebagai pemisah desimal (format Indonesia)."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v or "")
    if f == int(f):
        return str(int(f))
    return repr(f).replace(".", ",")


def build_id_actual(r: dict) -> str:
    """Id Actual = gabungan Kebun + Afdeling + Blok + Code LSU + Koord_X + Koord_Y (tanpa pemisah)."""
    return (f"{r.get('kebun','')}{r.get('afdeling','')}{r.get('blok','')}{r.get('code_lsu','')}"
            f"{fmt_num(r.get('koord_x'))}{fmt_num(r.get('koord_y'))}")


def compute_sph(r: dict) -> float:
    """SPH (Stand Per Hectare) = Jumlah Pokok / Luas (Ha). Dibulatkan 2 desimal."""
    try:
        luas = float(r.get("luas_ha") or 0)
        pokok = float(r.get("jumlah_pokok") or 0)
    except (TypeError, ValueError):
        return 0
    return round(pokok / luas, 2) if luas else 0


def _fnum(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def compute_la(r: dict) -> float:
    """Luas Daun (LA) per pokok, m2 — metode Hardon (1969) utk kelapa sawit.
    LA = 0.55 x jumlah_anak_daun x rata2_panjang_helai(cm) x rata2_lebar_helai(cm)
         x jumlah_pelepah / 10000.
    Rata-rata memakai nilai helai 1 & 2 yang terisi (>0).
    """
    lens = [x for x in (_fnum(r.get("panjang_helai_1")), _fnum(r.get("panjang_helai_2"))) if x > 0]
    wids = [x for x in (_fnum(r.get("lebar_helai_1")), _fnum(r.get("lebar_helai_2"))) if x > 0]
    n = _fnum(r.get("jumlah_anak_daun"))
    fronds = _fnum(r.get("jumlah_pelepah"))
    if not lens or not wids or n <= 0:
        return 0
    mean_l = sum(lens) / len(lens)
    mean_w = sum(wids) / len(wids)
    la_frond = 0.55 * n * mean_l * mean_w / 10000.0  # m2 per pelepah
    la = la_frond * fronds if fronds > 0 else la_frond
    return round(la, 4)


def compute_lai(r: dict) -> float:
    """LAI (Leaf Area Index) = LA per pokok (m2) x SPH / 10000 (SPH = pokok/ha)."""
    la = compute_la(r)
    sph = compute_sph(r)
    if la <= 0 or sph <= 0:
        return 0
    return round(la * sph / 10000.0, 4)


def apply_derived(d: dict) -> dict:
    """Isi field turunan otomatis: sph, la, lai."""
    d["sph"] = compute_sph(d)
    d["la"] = compute_la(d)
    d["lai"] = compute_lai(d)
    return d


def num_out(v) -> str:
    """Format angka umum untuk ekspor: buang desimal jika bulat, koma untuk desimal."""
    if v in (None, ""):
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if f == int(f):
        return str(int(f))
    return repr(round(f, 4)).replace(".", ",")


# Definisi kolom data (key, judul) — dipakai bersama untuk Excel & PDF
FIELD_COLUMNS = [
    ("id_actual", "ID Actual", "text"),
    ("kebun", "Kebun", "text"),
    ("afdeling", "Afdeling", "text"),
    ("code_lsu", "Kode LSU", "text"),
    ("blok", "Block", "text"),
    ("luas_ha", "Luas (Ha)", "num"),
    ("jumlah_pokok", "Jumlah Pokok", "num"),
    ("titik_sample", "Titik Sample", "text"),
    ("koord_x", "Koordinat (X)", "coord"),
    ("koord_y", "Koordinat (Y)", "coord"),
    ("kategori", "Kategori", "text"),
    ("keterangan", "Keterangan", "text"),
    ("sph", "SPH", "num"),
    ("jumlah_pelepah", "Jumlah pelepah", "num"),
    ("panjang_pelepah", "Panjang pelepah (cm)", "num"),
    ("lebar_petiol", "Lebar petiol (cm)", "num"),
    ("tebal_petiol", "Tebal petiol (cm)", "num"),
    ("panjang_helai_1", "Panjang helai anak daun 1 (cm)", "num"),
    ("panjang_helai_2", "Panjang helai anak daun 2 (cm)", "num"),
    ("lebar_helai_1", "Lebar helai anak daun 1 (cm)", "num"),
    ("lebar_helai_2", "Lebar helai anak daun 2 (cm)", "num"),
    ("jumlah_anak_daun", "Jumlah anak daun (helai)", "num"),
    ("tanggal_lsu", "Tanggal LSU", "text"),
    ("la", "LA", "num"),
    ("lai", "LAI", "num"),
]


def col_value(d: dict, key: str, kind: str) -> str:
    v = d.get(key, "")
    if kind == "coord":
        return fmt_num(v)
    if kind == "num":
        return num_out(v)
    return "" if v is None else str(v)


def build_payload(r: dict) -> str:
    """Isi QR = sama persis dengan Id Actual (nilai tergabung tanpa label/pemisah)."""
    return build_id_actual(r)


def make_qr_image(payload: str):
    qr = qrcode.QRCode(version=None, box_size=10, border=2,
                       error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(payload)
    qr.make(fit=True)
    return qr.make_image(fill_color="#0F291E", back_color="white").convert("RGB")


# ---------------------------------------------------------------- Auth routes
@api_router.post("/auth/register")
async def register(body: UserRegister, response: Response):
    email = body.email.strip().lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {"email": email, "password_hash": hash_password(body.password),
           "name": body.name, "role": "user", "created_at": datetime.now(timezone.utc).isoformat()}
    res = await db.users.insert_one(doc)
    token = create_access_token(str(res.inserted_id), email)
    set_auth_cookie(response, token)
    return {"id": str(res.inserted_id), "email": email, "name": body.name, "role": "user"}


@api_router.post("/auth/login")
async def login(body: UserLogin, response: Response):
    email = body.email.strip().lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    token = create_access_token(str(user["_id"]), email)
    set_auth_cookie(response, token)
    return {"id": str(user["_id"]), "email": email, "name": user.get("name"), "role": user.get("role")}


@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["_id"], "email": user["email"], "name": user.get("name"), "role": user.get("role")}


# ---------------------------------------------------------------- Records
async def next_id_actual() -> int:
    doc = await db.counters.find_one_and_update(
        {"_id": "id_actual"}, {"$inc": {"seq": 1}},
        upsert=True, return_document=True)
    return doc["seq"]


def serialize(doc: dict) -> dict:
    doc["_id"] = str(doc["_id"])
    doc["id_actual"] = build_id_actual(doc)
    doc["payload"] = build_payload(doc)
    apply_derived(doc)
    return doc


@api_router.get("/records")
async def list_records(user: dict = Depends(get_current_user)):
    docs = await db.records.find().sort("created_at", 1).to_list(5000)
    return [serialize(d) for d in docs]


@api_router.get("/records/stats")
async def stats(user: dict = Depends(get_current_user)):
    docs = await db.records.find().to_list(5000)
    kebun = set(); afd = set(); blok = set(); lsu = set()
    for d in docs:
        kebun.add(d.get("kebun", "")); afd.add((d.get("kebun",""), d.get("afdeling","")))
        blok.add(d.get("blok", "")); lsu.add(d.get("code_lsu", ""))
    return {"total_records": len(docs), "total_kebun": len([k for k in kebun if k]),
            "total_afdeling": len([a for a in afd if a[1]]),
            "total_blok": len([b for b in blok if b]),
            "total_lsu": len([l for l in lsu if l])}


@api_router.post("/records")
async def create_record(body: RecordInput, user: dict = Depends(get_current_user)):
    doc = body.model_dump()
    apply_derived(doc)
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    if _is_tagged(doc):
        doc["tagged_at"] = doc["created_at"]
    res = await db.records.insert_one(doc)
    saved = await db.records.find_one({"_id": res.inserted_id})
    return serialize(saved)


@api_router.put("/records/{rid}")
async def update_record(rid: str, body: RecordInput, user: dict = Depends(get_current_user)):
    existing = await db.records.find_one({"_id": ObjectId(rid)})
    if not existing:
        raise HTTPException(status_code=404, detail="Data tidak ditemukan")
    new_data = body.model_dump()
    apply_derived(new_data)
    update = {"$set": dict(new_data)}
    now_tagged = _is_tagged(new_data)
    if now_tagged and not existing.get("tagged_at"):
        update["$set"]["tagged_at"] = datetime.now(timezone.utc).isoformat()
    elif not now_tagged and existing.get("tagged_at"):
        update["$unset"] = {"tagged_at": ""}
    await db.records.update_one({"_id": ObjectId(rid)}, update)
    saved = await db.records.find_one({"_id": ObjectId(rid)})
    return serialize(saved)


@api_router.delete("/records/{rid}")
async def delete_record(rid: str, user: dict = Depends(get_current_user)):
    await db.records.delete_one({"_id": ObjectId(rid)})
    return {"ok": True}


@api_router.delete("/records")
async def delete_all(user: dict = Depends(get_current_user)):
    await db.records.delete_many({})
    return {"ok": True}


@api_router.post("/records/delete-bulk")
async def delete_bulk(body: IdList, user: dict = Depends(get_current_user)):
    if not body.ids:
        return {"ok": True, "deleted": 0}
    oids = [ObjectId(i) for i in body.ids]
    res = await db.records.delete_many({"_id": {"$in": oids}})
    return {"ok": True, "deleted": res.deleted_count}


# ---------------------------------------------------------------- Excel
@api_router.get("/records/template")
async def download_template(user: dict = Depends(get_current_user)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"
    headers = [
        "Kebun", "Afdeling", "Kode LSU", "Block", "Luas (Ha)", "Jumlah Pokok",
        "Titik Sample", "Koordinat (X)", "Koordinat (Y)", "Kategori", "Keterangan",
        "Jumlah pelepah", "Panjang pelepah (cm)", "Lebar petiol (cm)", "Tebal petiol (cm)",
        "Panjang helai anak daun 1 (cm)", "Panjang helai anak daun 2 (cm)",
        "Lebar helai anak daun 1 (cm)", "Lebar helai anak daun 2 (cm)",
        "Jumlah anak daun (helai)", "Tanggal LSU", "LA", "LAI",
    ]
    ws.append(headers)
    ws.append([
        "KSL", "OA", "TS01", "OA11", 4.5, 630, "TS-01", 110.400113, 0.654521,
        "Kategori A", "Contoh keterangan", 40, 550, 5.2, 3.1,
        120, 118, 6.5, 6.3, 250, "2025-07-01", 12.5, 3.2,
    ])
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_kebun.xlsx"})


def _cell(v):
    return "" if v is None else str(v).strip()


def _num(v):
    try:
        return float(str(v).replace(",", ".").strip())
    except (ValueError, AttributeError):
        return 0.0


def parse_excel(content: bytes) -> List[dict]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File Excel tidak valid")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong")

    def norm(s):
        return "".join(ch for ch in str(s or "").lower() if ch.isalnum())

    header = [norm(h) for h in rows[0]]

    # field -> (kind, [normalized aliases])
    FIELD_ALIASES = {
        "kebun": ("text", ["kebun"]),
        "afdeling": ("text", ["afdeling", "afd"]),
        "code_lsu": ("text", ["kodelsu", "codelsu", "lsu"]),
        "blok": ("text", ["block", "blok"]),
        "luas_ha": ("num", ["luasha", "luas"]),
        "jumlah_pokok": ("num", ["jumlahpokok", "pokok"]),
        "titik_sample": ("text", ["titiksample", "titiksampel"]),
        "koord_x": ("num", ["koordinatx", "koordx", "x"]),
        "koord_y": ("num", ["koordinaty", "koordy", "y"]),
        "kategori": ("text", ["kategori"]),
        "keterangan": ("text", ["keterangan"]),
        "jumlah_pelepah": ("num", ["jumlahpelepah"]),
        "panjang_pelepah": ("num", ["panjangpelepahcm", "panjangpelepah"]),
        "lebar_petiol": ("num", ["lebarpetiolcm", "lebarpetiol"]),
        "tebal_petiol": ("num", ["tebalpetiolcm", "tebalpetiol"]),
        "panjang_helai_1": ("num", ["panjanghelaianakdaun1cm", "panjanghelaianakdaun1", "panjanghelai1"]),
        "panjang_helai_2": ("num", ["panjanghelaianakdaun2cm", "panjanghelaianakdaun2", "panjanghelai2"]),
        "lebar_helai_1": ("num", ["lebarhelaianakdaun1cm", "lebarhelaianakdaun1", "lebarhelai1"]),
        "lebar_helai_2": ("num", ["lebarhelaianakdaun2cm", "lebarhelaianakdaun2", "lebarhelai2"]),
        "jumlah_anak_daun": ("num", ["jumlahanakdaunhelai", "jumlahanakdaun"]),
        "tanggal_lsu": ("text", ["tanggallsu"]),
        "la": ("num", ["la"]),
        "lai": ("num", ["lai"]),
    }

    def find_idx(aliases):
        for a in aliases:
            if a in header:
                return header.index(a)
        return None

    field_idx = {f: (kind, find_idx(al)) for f, (kind, al) in FIELD_ALIASES.items()}

    parsed = []
    for row in rows[1:]:
        if row is None or all(c is None or _cell(c) == "" for c in row):
            continue

        def g(i):
            return _cell(row[i]) if i is not None and i < len(row) else ""

        def gn(i):
            return _num(row[i]) if i is not None and i < len(row) else 0.0

        rec = {}
        for f, (kind, i) in field_idx.items():
            rec[f] = gn(i) if kind == "num" else g(i)
        if not (rec.get("kebun") or rec.get("blok") or rec.get("code_lsu")):
            continue
        parsed.append(rec)
    return parsed


@api_router.post("/records/import/preview")
async def import_preview(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    content = await file.read()
    rows = parse_excel(content)
    return {"rows": rows, "count": len(rows)}


@api_router.post("/records/import/confirm")
async def import_confirm(body: ImportConfirm, user: dict = Depends(get_current_user)):
    inserted = 0
    for r in body.rows:
        doc = r.model_dump()
        apply_derived(doc)
        doc["created_at"] = datetime.now(timezone.utc).isoformat()
        if _is_tagged(doc):
            doc["tagged_at"] = doc["created_at"]
        await db.records.insert_one(doc)
        inserted += 1
    return {"inserted": inserted}


# ---------------------------------------------------------------- PDF
def _draw_qr(c, payload, x, y, size):
    img = make_qr_image(payload)
    bio = io.BytesIO(); img.save(bio, format="PNG"); bio.seek(0)
    c.drawImage(ImageReader(bio), x, y, size, size)


async def _fetch_docs(ids: Optional[List[str]]):
    if ids:
        oids = [ObjectId(i) for i in ids]
        docs = await db.records.find({"_id": {"$in": oids}}).sort("created_at", 1).to_list(5000)
    else:
        docs = await db.records.find().sort("created_at", 1).to_list(5000)
    for d in docs:
        d["id_actual"] = build_id_actual(d)
        apply_derived(d)
    return docs


LABEL_LAYOUTS = {
    "small": (4, 6),   # label kecil, 24 per halaman
    "medium": (3, 5),  # sedang, 15 per halaman (default)
    "large": (2, 3),   # besar, 6 per halaman
}


def build_labels_pdf(docs, size: str = "medium") -> io.BytesIO:
    buf = io.BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=A4)
    pw, ph = A4
    cols, rows_pp = LABEL_LAYOUTS.get(size, LABEL_LAYOUTS["medium"])
    margin = 12 * mm
    cell_w = (pw - 2 * margin) / cols
    cell_h = (ph - 2 * margin) / rows_pp
    per_page = cols * rows_pp
    for i, d in enumerate(docs):
        pos = i % per_page
        if i > 0 and pos == 0:
            c.showPage()
        col = pos % cols
        rowi = pos // cols
        cx = margin + col * cell_w
        cy = ph - margin - (rowi + 1) * cell_h
        # Outer border box (solid, seperti contoh label)
        bx, by = cx + 3, cy + 3
        bw, bh = cell_w - 6, cell_h - 6
        c.setStrokeColor(colors.HexColor("#111827"))
        c.setLineWidth(1)
        c.rect(bx, by, bw, bh, stroke=1, fill=0)
        # Area teks di bawah (2 baris) dipisahkan garis horizontal
        text_area_h = {"small": 9 * mm, "medium": 12 * mm, "large": 16 * mm}.get(size, 12 * mm)
        f1 = {"small": 8, "medium": 9, "large": 13}.get(size, 9)
        f2 = {"small": 7, "medium": 8, "large": 11}.get(size, 8)
        divider_y = by + text_area_h
        c.setLineWidth(0.8)
        c.line(bx, divider_y, bx + bw, divider_y)
        # QR di area atas
        qr_size = min(bw - 8 * mm, bh - text_area_h - 8 * mm)
        qr_x = bx + (bw - qr_size) / 2
        qr_y = divider_y + ((bh - text_area_h) - qr_size) / 2
        _draw_qr(c, build_payload(d), qr_x, qr_y, qr_size)
        # Teks keterangan
        c.setFillColor(colors.HexColor("#111827"))
        line1 = " ".join(
            str(d.get(k, "")).strip()
            for k in ("kebun", "afdeling", "titik_sample", "blok")
            if str(d.get(k, "")).strip()
        )
        line2 = f"{fmt_num(d.get('koord_x'))} {fmt_num(d.get('koord_y'))}".strip()
        c.setFont("Helvetica-Bold", f1)
        c.drawCentredString(bx + bw / 2, divider_y - (f1 + 3), line1[:40])
        c.setFont("Helvetica-Bold", f2)
        c.drawCentredString(bx + bw / 2, by + 2.5 * mm, line2[:40])
    if not docs:
        c.setFont("Helvetica", 12)
        c.drawCentredString(pw / 2, ph / 2, "Tidak ada data")
    c.showPage(); c.save(); buf.seek(0)
    return buf


def _coord_filled(v) -> bool:
    s = str(v if v is not None else "").strip()
    if s == "":
        return False
    try:
        if float(s.replace(",", ".")) == 0:
            return False  # nilai 0 dianggap belum di-tagging
    except ValueError:
        pass
    return True


def _is_tagged(d) -> bool:
    return _coord_filled(d.get("koord_x")) and _coord_filled(d.get("koord_y"))


def build_table_pdf(docs) -> io.BytesIO:
    buf = io.BytesIO()
    page = landscape(A4)
    pw, ph = page
    margin = 12 * mm
    doc_tpl = SimpleDocTemplate(
        buf, pagesize=page,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin)

    hstyle = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=5.2,
                            leading=6, textColor=colors.white, alignment=TA_CENTER)
    cstyle = ParagraphStyle("c", fontName="Helvetica", fontSize=5.2,
                            leading=6, textColor=colors.HexColor("#1F2937"), alignment=TA_CENTER)
    title_style = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=14,
                                 textColor=colors.HexColor("#0F291E"))
    sub_style = ParagraphStyle("s", fontName="Helvetica", fontSize=8,
                               textColor=colors.HexColor("#4B5563"))

    columns = FIELD_COLUMNS + [("_status", "Keterangan Tagging", "text")]
    # bobot lebar kolom relatif
    weights = {
        "id_actual": 4.2, "kebun": 1.4, "afdeling": 1.2, "code_lsu": 1.4, "blok": 1.4,
        "luas_ha": 1.1, "jumlah_pokok": 1.3, "titik_sample": 1.4, "koord_x": 1.9,
        "koord_y": 1.9, "kategori": 1.6, "keterangan": 2.0, "sph": 1.0,
        "jumlah_pelepah": 1.3, "panjang_pelepah": 1.4, "lebar_petiol": 1.3,
        "tebal_petiol": 1.3, "panjang_helai_1": 1.5, "panjang_helai_2": 1.5,
        "lebar_helai_1": 1.4, "lebar_helai_2": 1.4, "jumlah_anak_daun": 1.5,
        "tanggal_lsu": 1.6, "la": 0.9, "lai": 0.9, "_status": 1.8,
    }
    usable = pw - 2 * margin
    total_w = sum(weights[k] for (k, _t, _knd) in columns)
    col_widths = [usable * weights[k] / total_w for (k, _t, _knd) in columns]

    header_row = [Paragraph(title, hstyle) for (_k, title, _t) in columns]
    data = [header_row]
    for d in docs:
        row = []
        for (k, _title, knd) in columns:
            if k == "_status":
                val = "Sudah di-tagging" if _is_tagged(d) else "Belum di-tagging"
            else:
                val = col_value(d, k, knd)
            row.append(Paragraph(str(val), cstyle))
        data.append(row)

    elements = [
        Paragraph("Laporan Data Kebun & Label QR", title_style),
        Spacer(1, 3 * mm),
        Paragraph(f"Total {len(docs)} data  -  {datetime.now().strftime('%d/%m/%Y %H:%M')}", sub_style),
        Spacer(1, 4 * mm),
    ]
    if not docs:
        elements.append(Paragraph("Tidak ada data", sub_style))
    else:
        tbl = Table(data, colWidths=col_widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F291E")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LEFTPADDING", (0, 0), (-1, -1), 1.5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 1.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ]))
        elements.append(tbl)
    doc_tpl.build(elements)
    buf.seek(0)
    return buf


def build_untagged_pdf(docs) -> io.BytesIO:
    """Daftar checklist lokasi yang BELUM di-tagging untuk petugas lapangan."""
    buf = io.BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=A4)
    pw, ph = A4
    margin = 14 * mm

    # kolom: No | Kebun | Afdeling | Blok | Code LSU | Titik Sample | Koord X (isi) | Koord Y (isi)
    col_x = [margin, margin + 12 * mm, margin + 40 * mm, margin + 60 * mm,
             margin + 82 * mm, margin + 108 * mm, margin + 140 * mm, margin + 172 * mm, pw - margin]
    headers = ["No", "Kebun", "Afdeling", "Blok", "Code LSU", "Titik Sample", "Koord X", "Koord Y"]
    row_h = 8 * mm

    def draw_header(y):
        c.setFillColor(colors.HexColor("#0F291E"))
        c.setFont("Helvetica-Bold", 15)
        c.drawString(margin, y, "Daftar Lokasi Belum di-tagging")
        y -= 7 * mm
        c.setFont("Helvetica", 9)
        c.setFillColor(colors.HexColor("#4B5563"))
        c.drawString(margin, y, f"Total {len(docs)} lokasi perlu di-tagging  -  {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        y -= 8 * mm
        # header row
        c.setFillColor(colors.HexColor("#0F291E"))
        c.rect(margin, y - row_h, pw - 2 * margin, row_h, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 9)
        for i, h in enumerate(headers):
            c.drawString(col_x[i] + 2 * mm, y - row_h + 2.6 * mm, h)
        return y - row_h

    y = draw_header(ph - margin)
    c.setFont("Helvetica", 9)
    for idx, d in enumerate(docs, 1):
        if y - row_h < margin:
            c.showPage()
            y = draw_header(ph - margin)
            c.setFont("Helvetica", 9)
        c.setStrokeColor(colors.HexColor("#E2E8F0"))
        c.setLineWidth(0.4)
        c.rect(margin, y - row_h, pw - 2 * margin, row_h, stroke=1, fill=0)
        for i in range(1, len(headers)):
            c.line(col_x[i], y - row_h, col_x[i], y)
        c.setFillColor(colors.HexColor("#1F2937"))
        vals = [str(idx), str(d.get("kebun", "")), str(d.get("afdeling", "")),
                str(d.get("blok", "")), str(d.get("code_lsu", "")),
                str(d.get("titik_sample", "")), "", ""]
        for i, v in enumerate(vals):
            c.drawString(col_x[i] + 2 * mm, y - row_h + 2.6 * mm, v[:20])
        y -= row_h
    if not docs:
        c.setFont("Helvetica", 12)
        c.setFillColor(colors.HexColor("#047857"))
        c.drawCentredString(pw / 2, ph / 2, "Semua lokasi sudah di-tagging.")
    c.showPage(); c.save(); buf.seek(0)
    return buf


def _pdf_response(buf, filename):
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})


@api_router.get("/records/export/labels")
async def export_labels(size: str = "medium", user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(None)
    return _pdf_response(build_labels_pdf(docs, size), "label_qr_kebun.pdf")


@api_router.post("/records/export/labels")
async def export_labels_selected(body: IdList, user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(body.ids)
    return _pdf_response(build_labels_pdf(docs, body.size), "label_qr_kebun.pdf")


@api_router.get("/records/export/table")
async def export_table(user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(None)
    return _pdf_response(build_table_pdf(docs), "laporan_tabel_kebun.pdf")


@api_router.post("/records/export/table")
async def export_table_selected(body: IdList, user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(body.ids)
    return _pdf_response(build_table_pdf(docs), "laporan_tabel_kebun.pdf")


@api_router.get("/records/export/untagged")
async def export_untagged(user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(None)
    untagged = [d for d in docs if not _is_tagged(d)]
    return _pdf_response(build_untagged_pdf(untagged), "daftar_belum_tagging.pdf")


def build_untagged_excel(docs) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Belum di-tagging"
    headers = ["No", "Kebun", "Afdeling", "Blok", "Code_LSU", "Titik Sample", "Luas (Ha)", "Jumlah Pokok", "Koord_X", "Koord_Y", "Keterangan"]
    ws.append(headers)
    for i, d in enumerate(docs, 1):
        ws.append([
            i, d.get("kebun", ""), d.get("afdeling", ""), d.get("blok", ""),
            d.get("code_lsu", ""), d.get("titik_sample", ""),
            num_out(d.get("luas_ha")), num_out(d.get("jumlah_pokok")),
            "", "", "Belum di-tagging",
        ])
    widths = [6, 16, 12, 12, 14, 16, 12, 14, 16, 16, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


@api_router.get("/records/tagging-progress")
async def tagging_progress(user: dict = Depends(get_current_user)):
    """Perkembangan jumlah lokasi ter-tagging dari waktu ke waktu (harian, kumulatif)."""
    docs = await db.records.find().to_list(20000)
    per_day = {}
    tagged_total = 0
    for d in docs:
        if not _is_tagged(d):
            continue
        tagged_total += 1
        ts = d.get("tagged_at") or d.get("created_at") or ""
        day = str(ts)[:10]
        if len(day) != 10:
            continue
        per_day[day] = per_day.get(day, 0) + 1
    series = []
    cumulative = 0
    for day in sorted(per_day.keys()):
        cumulative += per_day[day]
        series.append({"date": day, "count": per_day[day], "cumulative": cumulative})
    return {
        "series": series,
        "tagged_total": tagged_total,
        "untagged_total": len(docs) - tagged_total,
        "total": len(docs),
    }


def build_excel(docs) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Kebun"
    headers = [title for (_k, title, _t) in FIELD_COLUMNS] + ["QR (isi)", "Status Tagging"]
    ws.append(headers)
    for d in docs:
        row = [col_value(d, k, t) for (k, _title, t) in FIELD_COLUMNS]
        row.append(build_payload(d))
        row.append("Sudah di-tagging" if _is_tagged(d) else "Belum di-tagging")
        ws.append(row)
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
    ws.column_dimensions["A"].width = 30
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def _excel_response(buf, filename):
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})


@api_router.get("/records/export/excel")
async def export_excel(user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(None)
    return _excel_response(build_excel(docs), "data_kebun.xlsx")


@api_router.post("/records/export/excel")
async def export_excel_selected(body: IdList, user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(body.ids)
    return _excel_response(build_excel(docs), "data_kebun.xlsx")


@api_router.get("/records/export/untagged-excel")
async def export_untagged_excel(user: dict = Depends(get_current_user)):
    docs = await _fetch_docs(None)
    untagged = [d for d in docs if not _is_tagged(d)]
    return _excel_response(build_untagged_excel(untagged), "daftar_belum_tagging.xlsx")


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@kebun.id").lower()
    admin_pw = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if existing is None:
        await db.users.insert_one({
            "email": admin_email, "password_hash": hash_password(admin_pw),
            "name": "Administrator", "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat()})
        logger.info("Admin seeded")
    elif not verify_password(admin_pw, existing["password_hash"]):
        await db.users.update_one({"email": admin_email},
                                  {"$set": {"password_hash": hash_password(admin_pw)}})


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
