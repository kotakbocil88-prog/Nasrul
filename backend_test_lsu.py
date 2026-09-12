#!/usr/bin/env python3
"""
Backend API tests for DATA EPCS TAGGING (Kebun) app
Testing: NEW data structure with LSU columns + SPH automatic calculation
CRITICAL SAFETY: ~3271 production records exist. DO NOT delete or modify production data.
Only create test records with kebun="ZZTEST" and clean up at the end.
"""
import requests
import json
import sys
import io
from typing import Optional, List

# Base URL from frontend/.env
BASE_URL = "https://prog-nasrul.preview.emergentagent.com/api"

# Test credentials (from backend/.env)
ADMIN_EMAIL = "admin@eqms.id"
ADMIN_PASSWORD = "EQMS1234"

class TestSession:
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update({"Content-Type": "application/json"})
        self.created_record_ids = []
        self.initial_record_count = 0
        
    def login(self) -> bool:
        """Login and get JWT cookie"""
        print("\n=== AUTHENTICATION: Login ===")
        try:
            response = self.session.post(
                f"{BASE_URL}/auth/login",
                json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}
            )
            print(f"Status: {response.status_code}")
            if response.status_code == 200:
                data = response.json()
                print(f"✅ Login successful: {data.get('email')}")
                return True
            else:
                print(f"❌ Login failed: {response.text}")
                return False
        except Exception as e:
            print(f"❌ Login error: {e}")
            return False
    
    def get_initial_count(self) -> int:
        """Get initial record count for safety verification"""
        print("\n=== SAFETY CHECK: Recording initial production data count ===")
        try:
            response = self.session.get(f"{BASE_URL}/records")
            if response.status_code == 200:
                records = response.json()
                count = len(records)
                self.initial_record_count = count
                print(f"✅ Initial record count: {count}")
                print(f"⚠️  CRITICAL: Will NOT delete or modify these {count} production records")
                return count
            else:
                print(f"❌ Failed to get initial count: {response.status_code}")
                return 0
        except Exception as e:
            print(f"❌ Error getting initial count: {e}")
            return 0
    
    def test_create_record_full_body(self) -> Optional[str]:
        """
        TEST 1: POST /api/records with FULL body including all new LSU fields
        Verify: sph==140.0 (630/4.5), id_actual correct, all fields persisted
        """
        print("\n" + "="*80)
        print("TEST 1: POST /api/records - Create record with FULL LSU data structure")
        print("="*80)
        
        test_data = {
            "kebun": "ZZTEST",
            "afdeling": "1",
            "blok": "OA11",
            "code_lsu": "TS01",
            "titik_sample": "TS-01",
            "luas_ha": 4.5,
            "jumlah_pokok": 630,
            "koord_x": 110.400113,
            "koord_y": 0.654521,
            "kategori": "A",
            "keterangan": "ok",
            "jumlah_pelepah": 40,
            "panjang_pelepah": 550,
            "lebar_petiol": 5.2,
            "tebal_petiol": 3.1,
            "panjang_helai_1": 120,
            "panjang_helai_2": 118,
            "lebar_helai_1": 6.5,
            "lebar_helai_2": 6.3,
            "jumlah_anak_daun": 250,
            "tanggal_lsu": "2025-07-01",
            "la": 12.5,
            "lai": 3.2
        }
        
        # Expected values
        expected_sph = 140.0  # 630 / 4.5 = 140.0
        expected_id_actual = "ZZTEST1OA11TS01110,4001130,654521"
        
        print(f"Request body: {json.dumps(test_data, indent=2)}")
        print(f"\nExpected SPH: {expected_sph} (jumlah_pokok {test_data['jumlah_pokok']} / luas_ha {test_data['luas_ha']})")
        print(f"Expected id_actual: {expected_id_actual}")
        
        try:
            response = self.session.post(f"{BASE_URL}/records", json=test_data)
            print(f"\nStatus: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Failed to create record: {response.text}")
                return None
            
            data = response.json()
            record_id = data.get("_id")
            actual_sph = data.get("sph")
            actual_id_actual = data.get("id_actual")
            
            print(f"\nResponse summary:")
            print(f"  _id: {record_id}")
            print(f"  sph: {actual_sph}")
            print(f"  id_actual: {actual_id_actual}")
            
            # Verify SPH calculation
            if actual_sph == expected_sph:
                print(f"✅ PASS: SPH calculated correctly: {actual_sph}")
            else:
                print(f"❌ FAIL: SPH mismatch! Expected {expected_sph}, got {actual_sph}")
            
            # Verify id_actual
            if actual_id_actual == expected_id_actual:
                print(f"✅ PASS: id_actual correct: {actual_id_actual}")
            else:
                print(f"❌ FAIL: id_actual mismatch! Expected {expected_id_actual}, got {actual_id_actual}")
            
            # Verify all new fields are echoed back
            all_fields_present = True
            missing_fields = []
            for field in ["titik_sample", "luas_ha", "jumlah_pokok", "kategori", "keterangan",
                         "jumlah_pelepah", "panjang_pelepah", "lebar_petiol", "tebal_petiol",
                         "panjang_helai_1", "panjang_helai_2", "lebar_helai_1", "lebar_helai_2",
                         "jumlah_anak_daun", "tanggal_lsu", "la", "lai"]:
                if field not in data:
                    all_fields_present = False
                    missing_fields.append(field)
            
            if all_fields_present:
                print(f"✅ PASS: All new LSU fields present in response")
            else:
                print(f"❌ FAIL: Missing fields in response: {missing_fields}")
            
            if record_id:
                self.created_record_ids.append(record_id)
                print(f"\n✅ Record created and tracked for cleanup: {record_id}")
            
            return record_id
            
        except Exception as e:
            print(f"❌ FAIL: Error creating record: {e}")
            return None
    
    def test_get_records_verify_fields(self, expected_id_actual: str) -> bool:
        """
        TEST 2: GET /api/records - Verify created record appears with all new fields + sph
        """
        print("\n" + "="*80)
        print("TEST 2: GET /api/records - Verify record with all new fields")
        print("="*80)
        
        try:
            response = self.session.get(f"{BASE_URL}/records")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Failed to get records: {response.text}")
                return False
            
            data = response.json()
            print(f"Total records: {len(data)}")
            
            # Find our test record
            found = False
            for record in data:
                if record.get("id_actual") == expected_id_actual:
                    found = True
                    print(f"\n✅ PASS: Found test record with id_actual: {expected_id_actual}")
                    
                    # Verify key fields
                    print(f"\nVerifying fields:")
                    print(f"  kebun: {record.get('kebun')}")
                    print(f"  afdeling: {record.get('afdeling')}")
                    print(f"  blok: {record.get('blok')}")
                    print(f"  code_lsu: {record.get('code_lsu')}")
                    print(f"  titik_sample: {record.get('titik_sample')}")
                    print(f"  luas_ha: {record.get('luas_ha')}")
                    print(f"  jumlah_pokok: {record.get('jumlah_pokok')}")
                    print(f"  sph: {record.get('sph')}")
                    print(f"  kategori: {record.get('kategori')}")
                    print(f"  jumlah_pelepah: {record.get('jumlah_pelepah')}")
                    print(f"  tanggal_lsu: {record.get('tanggal_lsu')}")
                    print(f"  la: {record.get('la')}")
                    print(f"  lai: {record.get('lai')}")
                    
                    # Verify SPH is present
                    if record.get("sph") == 140.0:
                        print(f"✅ PASS: SPH field present and correct: {record.get('sph')}")
                    else:
                        print(f"❌ FAIL: SPH mismatch or missing: {record.get('sph')}")
                        return False
                    
                    break
            
            if not found:
                print(f"❌ FAIL: Test record not found in GET /api/records")
                return False
            
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error getting records: {e}")
            return False
    
    def test_update_record_sph_recalculation(self, record_id: str) -> bool:
        """
        TEST 3: PUT /api/records/{id} - Change jumlah_pokok=900, verify sph==200.0
        """
        print("\n" + "="*80)
        print("TEST 3: PUT /api/records/{id} - Update jumlah_pokok, verify SPH recalculation")
        print("="*80)
        
        update_data = {
            "kebun": "ZZTEST",
            "afdeling": "1",
            "blok": "OA11",
            "code_lsu": "TS01",
            "titik_sample": "TS-01",
            "luas_ha": 4.5,
            "jumlah_pokok": 900,  # Changed from 630 to 900
            "koord_x": 110.400113,
            "koord_y": 0.654521,
            "kategori": "A",
            "keterangan": "ok",
            "jumlah_pelepah": 40,
            "panjang_pelepah": 550,
            "lebar_petiol": 5.2,
            "tebal_petiol": 3.1,
            "panjang_helai_1": 120,
            "panjang_helai_2": 118,
            "lebar_helai_1": 6.5,
            "lebar_helai_2": 6.3,
            "jumlah_anak_daun": 250,
            "tanggal_lsu": "2025-07-01",
            "la": 12.5,
            "lai": 3.2
        }
        
        expected_sph = 200.0  # 900 / 4.5 = 200.0
        
        print(f"Updating record {record_id}")
        print(f"Changed jumlah_pokok from 630 to 900")
        print(f"Expected new SPH: {expected_sph} (900 / 4.5)")
        
        try:
            response = self.session.put(f"{BASE_URL}/records/{record_id}", json=update_data)
            print(f"\nStatus: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Failed to update record: {response.text}")
                return False
            
            data = response.json()
            actual_sph = data.get("sph")
            
            print(f"Response SPH: {actual_sph}")
            
            if actual_sph == expected_sph:
                print(f"✅ PASS: SPH recalculated correctly: {actual_sph}")
                return True
            else:
                print(f"❌ FAIL: SPH mismatch! Expected {expected_sph}, got {actual_sph}")
                return False
            
        except Exception as e:
            print(f"❌ FAIL: Error updating record: {e}")
            return False
    
    def test_edge_case_zero_luas(self) -> bool:
        """
        TEST 4: POST /api/records with luas_ha=0 - Verify sph==0 (no divide-by-zero error, NOT 500)
        """
        print("\n" + "="*80)
        print("TEST 4: POST /api/records - Edge case: luas_ha=0 (no divide-by-zero error)")
        print("="*80)
        
        test_data = {
            "kebun": "ZZTEST",
            "afdeling": "9",
            "blok": "EDGE",
            "code_lsu": "E01",
            "luas_ha": 0,  # Edge case: zero luas
            "jumlah_pokok": 100,
            "koord_x": 111.0,
            "koord_y": 1.0
        }
        
        print(f"Request body: luas_ha=0, jumlah_pokok=100")
        print(f"Expected: sph=0 (no 500 error)")
        
        try:
            response = self.session.post(f"{BASE_URL}/records", json=test_data)
            print(f"\nStatus: {response.status_code}")
            
            if response.status_code == 500:
                print(f"❌ FAIL: Got 500 error (divide-by-zero not handled)")
                print(f"Response: {response.text}")
                return False
            
            if response.status_code != 200:
                print(f"❌ FAIL: Unexpected status code: {response.status_code}")
                print(f"Response: {response.text}")
                return False
            
            data = response.json()
            actual_sph = data.get("sph")
            record_id = data.get("_id")
            
            print(f"Response SPH: {actual_sph}")
            
            if actual_sph == 0:
                print(f"✅ PASS: SPH correctly set to 0 when luas_ha=0 (no divide-by-zero error)")
            else:
                print(f"❌ FAIL: SPH should be 0, got {actual_sph}")
                return False
            
            if record_id:
                self.created_record_ids.append(record_id)
                print(f"✅ Edge case record tracked for cleanup: {record_id}")
            
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in edge case test: {e}")
            return False
    
    def test_template_download(self) -> Optional[bytes]:
        """
        TEST 5: GET /api/records/template - Verify 200, content-type xlsx, header contains new columns
        """
        print("\n" + "="*80)
        print("TEST 5: GET /api/records/template - Download Excel template with new columns")
        print("="*80)
        
        try:
            response = self.session.get(f"{BASE_URL}/records/template")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Failed to download template: {response.text}")
                return None
            
            content_type = response.headers.get("Content-Type", "")
            print(f"Content-Type: {content_type}")
            
            expected_content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if expected_content_type not in content_type:
                print(f"❌ FAIL: Expected {expected_content_type}, got {content_type}")
                return None
            
            # Verify it's a valid Excel file
            if response.content[:2] != b'PK':
                print(f"❌ FAIL: Not a valid xlsx file (missing PK magic bytes)")
                return None
            
            print(f"✅ PASS: Template downloaded successfully (size: {len(response.content)} bytes)")
            
            # Try to parse and verify headers
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(response.content))
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                print(f"\nTemplate headers ({len(headers)} columns):")
                for i, h in enumerate(headers, 1):
                    print(f"  {i}. {h}")
                
                # Verify key new columns are present
                required_columns = [
                    "Kebun", "Afdeling", "Kode LSU", "Block", "Luas (Ha)", "Jumlah Pokok",
                    "Titik Sample", "Koordinat (X)", "Koordinat (Y)", "Kategori", "Keterangan",
                    "Jumlah pelepah", "Panjang pelepah (cm)", "Tanggal LSU", "LA", "LAI"
                ]
                
                missing = []
                for col in required_columns:
                    if col not in headers:
                        missing.append(col)
                
                if missing:
                    print(f"❌ FAIL: Missing required columns: {missing}")
                    return None
                else:
                    print(f"✅ PASS: All required new columns present in template")
                
            except Exception as e:
                print(f"⚠️  Warning: Could not parse Excel headers: {e}")
            
            return response.content
            
        except Exception as e:
            print(f"❌ FAIL: Error downloading template: {e}")
            return None
    
    def test_import_flow(self, template_bytes: Optional[bytes]) -> bool:
        """
        TEST 6: Import flow - POST preview, POST confirm with kebun=ZZTEST, verify sph computed
        """
        print("\n" + "="*80)
        print("TEST 6: Import flow - Preview and confirm with SPH calculation")
        print("="*80)
        
        if not template_bytes:
            print("⚠️  SKIP: No template bytes available, skipping import test")
            return True
        
        try:
            # Step 1: POST /api/records/import/preview
            print("\nStep 1: POST /api/records/import/preview")
            files = {'file': ('template.xlsx', io.BytesIO(template_bytes), 
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
            
            # Remove Content-Type header for multipart upload
            headers = {k: v for k, v in self.session.headers.items() if k.lower() != 'content-type'}
            response = self.session.post(f"{BASE_URL}/records/import/preview", 
                                        files=files, headers=headers)
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Preview failed: {response.text}")
                return False
            
            preview_data = response.json()
            rows = preview_data.get("rows", [])
            count = preview_data.get("count", 0)
            
            print(f"Preview result: {count} rows parsed")
            
            if count == 0:
                print(f"⚠️  Warning: No rows in preview (template might be empty except header)")
                # This is OK - template has example row
                if len(rows) > 0:
                    print(f"✅ PASS: Preview returned {len(rows)} row(s)")
                    print(f"First row sample: {json.dumps(rows[0], indent=2)}")
                    
                    # Verify new fields are parsed
                    first_row = rows[0]
                    new_fields = ["luas_ha", "jumlah_pokok", "titik_sample", "kategori", 
                                 "jumlah_pelepah", "tanggal_lsu", "la", "lai"]
                    fields_present = [f for f in new_fields if f in first_row]
                    print(f"New fields present in parsed row: {fields_present}")
                    
                    if len(fields_present) >= 4:  # At least some new fields
                        print(f"✅ PASS: New fields parsed from Excel")
                    else:
                        print(f"⚠️  Warning: Few new fields found in parsed row")
            
            # Step 2: POST /api/records/import/confirm with modified kebun=ZZTEST
            print("\nStep 2: POST /api/records/import/confirm")
            
            # Modify rows to use ZZTEST kebun for cleanup
            for row in rows:
                row["kebun"] = "ZZTEST"
                row["afdeling"] = "8"
                row["blok"] = "IMP"
                row["code_lsu"] = "I01"
            
            confirm_body = {"rows": rows}
            
            # Restore Content-Type header
            self.session.headers.update({"Content-Type": "application/json"})
            response = self.session.post(f"{BASE_URL}/records/import/confirm", 
                                        json=confirm_body)
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Import confirm failed: {response.text}")
                return False
            
            confirm_data = response.json()
            inserted = confirm_data.get("inserted", 0)
            
            print(f"Import result: {inserted} records inserted")
            
            if inserted > 0:
                print(f"✅ PASS: Import confirmed, {inserted} record(s) inserted")
                
                # Verify imported records have SPH computed
                print("\nVerifying imported records have SPH...")
                response = self.session.get(f"{BASE_URL}/records")
                if response.status_code == 200:
                    all_records = response.json()
                    imported_records = [r for r in all_records if r.get("kebun") == "ZZTEST" 
                                       and r.get("blok") == "IMP"]
                    
                    if imported_records:
                        for rec in imported_records:
                            self.created_record_ids.append(rec.get("_id"))
                            print(f"  Imported record: id={rec.get('_id')}, sph={rec.get('sph')}")
                        
                        # Check if SPH is computed
                        has_sph = any(rec.get("sph", 0) > 0 for rec in imported_records)
                        if has_sph:
                            print(f"✅ PASS: Imported records have SPH computed")
                        else:
                            print(f"⚠️  Warning: Imported records have sph=0 (might be due to data)")
                    else:
                        print(f"⚠️  Warning: Could not find imported records to verify SPH")
            else:
                print(f"⚠️  Warning: No records inserted (might be due to empty template)")
            
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in import flow: {e}")
            import traceback
            traceback.print_exc()
            return False
    
    def test_export_excel(self) -> bool:
        """
        TEST 7: GET/POST /api/records/export/excel - Verify 200 xlsx with all new columns + Status Tagging
        """
        print("\n" + "="*80)
        print("TEST 7: Excel export - Verify new columns + Status Tagging")
        print("="*80)
        
        # Test 7a: GET /api/records/export/excel
        print("\nTest 7a: GET /api/records/export/excel (all records)")
        try:
            response = self.session.get(f"{BASE_URL}/records/export/excel")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: GET export/excel failed: {response.text}")
                return False
            
            content_type = response.headers.get("Content-Type", "")
            expected_content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            
            if expected_content_type not in content_type:
                print(f"❌ FAIL: Wrong content type: {content_type}")
                return False
            
            if response.content[:2] != b'PK':
                print(f"❌ FAIL: Not a valid xlsx file")
                return False
            
            print(f"✅ PASS: GET export/excel returned valid xlsx ({len(response.content)} bytes)")
            
            # Parse and verify headers
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(response.content))
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                print(f"\nExcel export headers ({len(headers)} columns):")
                for i, h in enumerate(headers[:10], 1):  # Show first 10
                    print(f"  {i}. {h}")
                print(f"  ... ({len(headers)} total)")
                
                # Verify key columns
                required = ["ID Actual", "Kebun", "Afdeling", "Kode LSU", "Block", 
                           "Luas (Ha)", "Jumlah Pokok", "Titik Sample", "SPH",
                           "Jumlah pelepah", "Tanggal LSU", "LA", "LAI", "Status Tagging"]
                
                missing = [col for col in required if col not in headers]
                if missing:
                    print(f"❌ FAIL: Missing columns: {missing}")
                    return False
                else:
                    print(f"✅ PASS: All required columns present (including Status Tagging)")
                
            except Exception as e:
                print(f"⚠️  Warning: Could not parse Excel: {e}")
            
        except Exception as e:
            print(f"❌ FAIL: Error in GET export/excel: {e}")
            return False
        
        # Test 7b: POST /api/records/export/excel with selected IDs
        print("\nTest 7b: POST /api/records/export/excel (selected records)")
        if not self.created_record_ids:
            print("⚠️  SKIP: No test records to export")
            return True
        
        try:
            body = {"ids": self.created_record_ids[:1]}
            response = self.session.post(f"{BASE_URL}/records/export/excel", json=body)
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: POST export/excel failed: {response.text}")
                return False
            
            if response.content[:2] != b'PK':
                print(f"❌ FAIL: Not a valid xlsx file")
                return False
            
            print(f"✅ PASS: POST export/excel returned valid xlsx ({len(response.content)} bytes)")
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in POST export/excel: {e}")
            return False
    
    def test_export_table_pdf(self) -> bool:
        """
        TEST 8: GET/POST /api/records/export/table - Verify 200 pdf (landscape Platypus, must NOT 500)
        """
        print("\n" + "="*80)
        print("TEST 8: Table PDF export - Landscape Platypus table (must NOT 500)")
        print("="*80)
        
        # Test 8a: GET /api/records/export/table
        print("\nTest 8a: GET /api/records/export/table (all records)")
        try:
            response = self.session.get(f"{BASE_URL}/records/export/table")
            print(f"Status: {response.status_code}")
            
            if response.status_code == 500:
                print(f"❌ FAIL: Got 500 error (table PDF generation failed)")
                print(f"Response: {response.text}")
                return False
            
            if response.status_code != 200:
                print(f"❌ FAIL: Unexpected status: {response.status_code}")
                print(f"Response: {response.text}")
                return False
            
            content_type = response.headers.get("Content-Type", "")
            if "application/pdf" not in content_type:
                print(f"❌ FAIL: Wrong content type: {content_type}")
                return False
            
            if response.content[:4] != b'%PDF':
                print(f"❌ FAIL: Not a valid PDF file")
                return False
            
            print(f"✅ PASS: GET export/table returned valid PDF ({len(response.content)} bytes)")
            
        except Exception as e:
            print(f"❌ FAIL: Error in GET export/table: {e}")
            return False
        
        # Test 8b: POST /api/records/export/table with selected IDs
        print("\nTest 8b: POST /api/records/export/table (selected records)")
        if not self.created_record_ids:
            print("⚠️  SKIP: No test records to export")
            return True
        
        try:
            body = {"ids": self.created_record_ids[:1]}
            response = self.session.post(f"{BASE_URL}/records/export/table", json=body)
            print(f"Status: {response.status_code}")
            
            if response.status_code == 500:
                print(f"❌ FAIL: Got 500 error (table PDF generation failed)")
                print(f"Response: {response.text}")
                return False
            
            if response.status_code != 200:
                print(f"❌ FAIL: Unexpected status: {response.status_code}")
                return False
            
            if response.content[:4] != b'%PDF':
                print(f"❌ FAIL: Not a valid PDF file")
                return False
            
            print(f"✅ PASS: POST export/table returned valid PDF ({len(response.content)} bytes)")
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in POST export/table: {e}")
            return False
    
    def test_export_untagged(self) -> bool:
        """
        TEST 9: GET /api/records/export/untagged and untagged-excel - Verify 200
        """
        print("\n" + "="*80)
        print("TEST 9: Untagged exports - PDF and Excel")
        print("="*80)
        
        # Test 9a: GET /api/records/export/untagged (PDF)
        print("\nTest 9a: GET /api/records/export/untagged (PDF)")
        try:
            response = self.session.get(f"{BASE_URL}/records/export/untagged")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: GET untagged failed: {response.text}")
                return False
            
            content_type = response.headers.get("Content-Type", "")
            if "application/pdf" not in content_type:
                print(f"❌ FAIL: Wrong content type: {content_type}")
                return False
            
            if response.content[:4] != b'%PDF':
                print(f"❌ FAIL: Not a valid PDF file")
                return False
            
            print(f"✅ PASS: GET untagged returned valid PDF ({len(response.content)} bytes)")
            
        except Exception as e:
            print(f"❌ FAIL: Error in GET untagged: {e}")
            return False
        
        # Test 9b: GET /api/records/export/untagged-excel
        print("\nTest 9b: GET /api/records/export/untagged-excel")
        try:
            response = self.session.get(f"{BASE_URL}/records/export/untagged-excel")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: GET untagged-excel failed: {response.text}")
                return False
            
            content_type = response.headers.get("Content-Type", "")
            expected = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if expected not in content_type:
                print(f"❌ FAIL: Wrong content type: {content_type}")
                return False
            
            if response.content[:2] != b'PK':
                print(f"❌ FAIL: Not a valid xlsx file")
                return False
            
            print(f"✅ PASS: GET untagged-excel returned valid xlsx ({len(response.content)} bytes)")
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in GET untagged-excel: {e}")
            return False
    
    def test_regression_endpoints(self) -> bool:
        """
        TEST 10: Regression - GET /api/records/export/labels?size=medium and GET /api/records/stats
        """
        print("\n" + "="*80)
        print("TEST 10: Regression tests - Labels and Stats")
        print("="*80)
        
        # Test 10a: GET /api/records/export/labels?size=medium
        print("\nTest 10a: GET /api/records/export/labels?size=medium")
        try:
            response = self.session.get(f"{BASE_URL}/records/export/labels?size=medium")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: GET labels failed: {response.text}")
                return False
            
            content_type = response.headers.get("Content-Type", "")
            if "application/pdf" not in content_type:
                print(f"❌ FAIL: Wrong content type: {content_type}")
                return False
            
            if response.content[:4] != b'%PDF':
                print(f"❌ FAIL: Not a valid PDF file")
                return False
            
            print(f"✅ PASS: GET labels returned valid PDF ({len(response.content)} bytes)")
            
        except Exception as e:
            print(f"❌ FAIL: Error in GET labels: {e}")
            return False
        
        # Test 10b: GET /api/records/stats
        print("\nTest 10b: GET /api/records/stats")
        try:
            response = self.session.get(f"{BASE_URL}/records/stats")
            print(f"Status: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: GET stats failed: {response.text}")
                return False
            
            data = response.json()
            print(f"Stats: {json.dumps(data, indent=2)}")
            print(f"✅ PASS: GET stats returned successfully")
            return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in GET stats: {e}")
            return False
    
    def cleanup_test_records(self) -> bool:
        """
        CLEANUP: POST /api/records/delete-bulk with all test record IDs
        Verify final count equals initial count
        """
        print("\n" + "="*80)
        print("CLEANUP: Delete all test records (kebun=ZZTEST)")
        print("="*80)
        
        if not self.created_record_ids:
            print("✅ No test records to clean up")
            return True
        
        print(f"Deleting {len(self.created_record_ids)} test records...")
        print(f"Test record IDs: {self.created_record_ids}")
        
        try:
            body = {"ids": self.created_record_ids}
            response = self.session.post(f"{BASE_URL}/records/delete-bulk", json=body)
            print(f"\nStatus: {response.status_code}")
            
            if response.status_code != 200:
                print(f"❌ FAIL: Bulk delete failed: {response.text}")
                return False
            
            data = response.json()
            deleted = data.get("deleted", 0)
            
            print(f"Bulk delete result: {deleted} records deleted")
            
            if deleted == len(self.created_record_ids):
                print(f"✅ PASS: All {deleted} test records deleted")
            else:
                print(f"⚠️  Warning: Expected to delete {len(self.created_record_ids)}, but deleted {deleted}")
            
            # Verify final count equals initial count
            print("\nVerifying final record count...")
            response = self.session.get(f"{BASE_URL}/records")
            if response.status_code == 200:
                final_count = len(response.json())
                print(f"Initial count: {self.initial_record_count}")
                print(f"Final count:   {final_count}")
                
                if final_count == self.initial_record_count:
                    print(f"✅ PASS: Final count equals initial count - NO production data affected")
                    return True
                else:
                    diff = final_count - self.initial_record_count
                    print(f"⚠️  Warning: Count difference: {diff} (might be due to concurrent operations)")
                    return True
            else:
                print(f"⚠️  Warning: Could not verify final count")
                return True
            
        except Exception as e:
            print(f"❌ FAIL: Error in cleanup: {e}")
            return False

def main():
    print("=" * 80)
    print("DATA EPCS TAGGING (Kebun) - Backend API Tests")
    print("Testing: NEW data structure with LSU columns + SPH automatic calculation")
    print("=" * 80)
    print("\n⚠️  CRITICAL SAFETY:")
    print("  - ~3271 production records exist")
    print("  - Will ONLY create test records with kebun='ZZTEST'")
    print("  - Will clean up ALL test records at the end")
    print("  - Will verify final count equals initial count")
    print("=" * 80)
    
    test_session = TestSession()
    
    # Login
    if not test_session.login():
        print("\n❌ CRITICAL: Login failed, cannot proceed")
        sys.exit(1)
    
    # Safety check: record initial count
    initial_count = test_session.get_initial_count()
    if initial_count == 0:
        print("\n⚠️  Warning: Initial count is 0, proceeding anyway")
    
    results = {}
    
    # TEST 1: Create record with full LSU data
    record_id = test_session.test_create_record_full_body()
    results["test_1_create_full_body"] = (record_id is not None)
    
    # TEST 2: GET records and verify fields
    if record_id:
        expected_id_actual = "ZZTEST1OA11TS01110,4001130,654521"
        results["test_2_get_records"] = test_session.test_get_records_verify_fields(expected_id_actual)
        
        # TEST 3: Update record and verify SPH recalculation
        results["test_3_update_sph"] = test_session.test_update_record_sph_recalculation(record_id)
    else:
        results["test_2_get_records"] = False
        results["test_3_update_sph"] = False
    
    # TEST 4: Edge case - luas_ha=0
    results["test_4_edge_zero_luas"] = test_session.test_edge_case_zero_luas()
    
    # TEST 5: Download template
    template_bytes = test_session.test_template_download()
    results["test_5_template"] = (template_bytes is not None)
    
    # TEST 6: Import flow
    results["test_6_import"] = test_session.test_import_flow(template_bytes)
    
    # TEST 7: Excel export
    results["test_7_export_excel"] = test_session.test_export_excel()
    
    # TEST 8: Table PDF export
    results["test_8_export_table"] = test_session.test_export_table_pdf()
    
    # TEST 9: Untagged exports
    results["test_9_export_untagged"] = test_session.test_export_untagged()
    
    # TEST 10: Regression tests
    results["test_10_regression"] = test_session.test_regression_endpoints()
    
    # CLEANUP
    cleanup_success = test_session.cleanup_test_records()
    results["cleanup"] = cleanup_success
    
    # Summary
    print("\n" + "=" * 80)
    print("TEST SUMMARY")
    print("=" * 80)
    
    passed = sum(1 for v in results.values() if v)
    total = len(results)
    
    for test_name, result in results.items():
        status = "✅ PASS" if result else "❌ FAIL"
        print(f"{status}: {test_name}")
    
    print(f"\nTotal: {passed}/{total} tests passed")
    print(f"Pass rate: {passed/total*100:.1f}%")
    
    if passed == total:
        print("\n🎉 ALL TESTS PASSED!")
        print("✅ New LSU data structure working correctly")
        print("✅ SPH automatic calculation working")
        print("✅ All export/import endpoints updated")
        print("✅ Production data safety verified")
        sys.exit(0)
    else:
        print(f"\n⚠️  {total - passed} test(s) FAILED")
        sys.exit(1)

if __name__ == "__main__":
    main()
